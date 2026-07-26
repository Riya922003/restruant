from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

HEADERS = [
    "Date", "Supplier", "Invoice #", "Category", "Subtotal", "Tax", "Total", "Payment Status",
]
_MONEY_COLS = (5, 6, 7)  # Subtotal, Tax, Total
_MONEY_FMT = "#,##0.00"


def build_workbook(rows: list[dict]) -> bytes:
    """Build the Expense Register .xlsx from approved invoice rows (spec 03 s11).
    Money as numbers (2dp), dates as dates, plus a totals row."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Expense Register"

    for col, header in enumerate(HEADERS, start=1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = Font(bold=True)

    total_sum = 0.0
    for r_index, row in enumerate(rows, start=2):
        ws.cell(row=r_index, column=1, value=row.get("invoice_date"))
        ws.cell(row=r_index, column=2, value=row.get("supplier_name"))
        ws.cell(row=r_index, column=3, value=row.get("invoice_number"))
        ws.cell(row=r_index, column=4, value=row.get("category_name") or "-")
        ws.cell(row=r_index, column=5, value=float(row.get("subtotal") or 0))
        ws.cell(row=r_index, column=6, value=float(row.get("tax") or 0))
        ws.cell(row=r_index, column=7, value=float(row.get("total") or 0))
        ws.cell(row=r_index, column=8, value=row.get("status"))
        total_sum += float(row.get("total") or 0)

    totals_row = len(rows) + 2
    label = ws.cell(row=totals_row, column=4, value="TOTAL")
    label.font = Font(bold=True)
    total_cell = ws.cell(row=totals_row, column=7, value=round(total_sum, 2))
    total_cell.font = Font(bold=True)

    for col in _MONEY_COLS:
        for r_index in range(2, totals_row + 1):
            ws.cell(row=r_index, column=col).number_format = _MONEY_FMT

    widths = [12, 28, 18, 20, 12, 10, 12, 16]
    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
